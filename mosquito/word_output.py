# -*- coding: utf-8 -*-
"""Word 输出：日报（叙述版，3.6/3.6.1）。

排版按示例文档实测；字体：中文统一 仿宋_GB2312，英文 Times New Roman。
（原“一览表 Word”输出已按需求取消，仅保留日报。）
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import config as C
from .pipeline import round1

_JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
_CENTER = WD_ALIGN_PARAGRAPH.CENTER
_NAN_TEXT = 'nan'


def _set_run(r, east, size, bold=None):
    r.font.name = C.FONT_EN
    rPr = r._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn('w:eastAsia'), east)
    r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold


def _shade_run(r, fill):
    """给 run 设置文字背景色（w:shd，填充色 fill 如 'FF0000'）"""
    rPr = r._element.get_or_add_rPr()
    shd = rPr.find(qn('w:shd'))
    if shd is None:
        shd = rPr.makeelement(qn('w:shd'), {})
        rPr.append(shd)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)


def _add_par(doc, text, size, align=None, indent_chars=0, line_pt=None, bold=None, keep_next=False):
    """添加段落；文本中出现 "nan" 字样时，对该文字单独设置红色背景（其余文字正常）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line_pt:
        pf.line_spacing = Pt(line_pt)          # 固定行距（exact）
    if keep_next:
        pf.keep_with_next = True
    if indent_chars:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        ind.set(qn('w:firstLine'), str(int(indent_chars * size * 20)))
    if text:
        parts = text.split(_NAN_TEXT)
        for i, part in enumerate(parts):
            if i > 0:                          # 每段 "nan" 前的分隔点 → 写入红色 "nan"
                r = p.add_run(_NAN_TEXT)
                _set_run(r, C.FONT_EN, size, bold)
                _shade_run(r, 'FF0000')
            if part:
                r = p.add_run(part)
                _set_run(r, C.FONT_CN, size, bold)
    return p


def _setup_page(doc, top, bottom, left, right):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(top)
    sec.bottom_margin = Cm(bottom)
    sec.left_margin = Cm(left)
    sec.right_margin = Cm(right)


# ---------------- 日报（3.6） ----------------

def build_section(final_df, exclude_field, excluded_cities, metric):
    """从最终表构建 Word 章节统计（metric = 'BI' 或 'ADI'）"""
    df = final_df
    value_col = 'BI*' if metric == 'BI' else 'ADI'

    # 地市列表（固定顺序、名称不带市、最后一个保留市、排除注记）
    cities = [c for c in C.CITY_ORDER if c in set(df['地市'])]
    if not cities:
        city_list = ''
    elif len(cities) == 1:
        c0 = cities[0]
        note = '（%s除外）' % exclude_field if exclude_field and c0 in excluded_cities else ''
        city_list = c0 + note + '市'
    else:
        names = []
        for c in cities:
            nm = c
            if exclude_field and c in excluded_cities:
                nm = '%s（%s除外）' % (c, exclude_field)
            names.append(nm)
        last = cities[-1]
        if exclude_field and last in excluded_cities:
            last_name = '%s市（%s除外）' % (last, exclude_field)
        else:
            last_name = names[-1] + '市'
        city_list = '、'.join(names[:-1]) + '和' + last_name

    total = len(df)
    district_num = df['区县'].nunique()                       # 区县名去重（市辖区合并，与参考一致）
    street_num = df.groupby(['地市', '街道']).ngroups
    point_num = df.groupby(['街道', '监测地点']).ngroups      # 监测点=（街道+监测地点）去重
    if metric == 'BI':
        ok_num = int((df[value_col] < 5).sum())
        risk_num = int((df[value_col] >= 5).sum())
    else:
        ok_num = int((df[value_col] <= 2).sum())
        risk_num = int((df[value_col] > 2).sum())
    ok_rate = round1(ok_num * 100.0 / total) if total else 0.0
    risk_rate = round1(risk_num * 100.0 / total) if total else 0.0

    # 风险列表（高→中→低，组内按值降序；数量为 0 不输出）
    # D63：元素格式 {地市+市}{区县}{街道}{监测地点}（{BI值}，第{N}天），N=监测天数
    def element(row):
        dist = row['区县'] if row['区县'] != '市辖区' else ''
        base = '%s市%s%s%s' % (row['地市'], dist, row['街道'], row['监测地点'])
        d = row.get('_days')
        if d is not None and not (isinstance(d, float) and d != d) and str(d).strip() != '':
            return '%s（%.1f，第%d天）' % (base, round1(row[value_col]), int(round(float(d))))
        return '%s（%.1f）' % (base, round1(row[value_col]))   # 无监测天数（如飞行监测）保持原格式

    if metric == 'BI':
        high = df[df[value_col] >= 20]
        mid = df[(df[value_col] >= 10) & (df[value_col] < 20)]
        low = df[(df[value_col] >= 5) & (df[value_col] < 10)]
    else:
        high = df[df[value_col] > 10]
        mid = df[(df[value_col] > 5) & (df[value_col] <= 10)]
        low = df[(df[value_col] > 2) & (df[value_col] <= 5)]

    risk_lists = {
        '高风险': [element(r) for _, r in high.sort_values(value_col, ascending=False).iterrows()],
        '中风险': [element(r) for _, r in mid.sort_values(value_col, ascending=False).iterrows()],
        '低风险': [element(r) for _, r in low.sort_values(value_col, ascending=False).iterrows()],
    }

    return {
        'metric': metric, 'city_list': city_list, 'city_num': len(cities),
        'district_num': district_num, 'street_num': street_num, 'point_num': point_num,
        'total': total, 'ok_num': ok_num, 'risk_num': risk_num,
        'ok_rate': ok_rate, 'risk_rate': risk_rate, 'risk_lists': risk_lists,
    }


def _write_section(doc, heading, sec, m, d):
    _add_par(doc, heading, 16, indent_chars=2, line_pt=28, keep_next=True)
    metric = sec['metric']
    if sec['total'] == 0:
        _add_par(doc, '%d月%d日，当日无符合条件（距末例天数≤5或>40000、防控区类型为核心区/警戒区）的%s监测数据。'
                 % (m, d, metric), 16, _JUSTIFY, indent_chars=2, line_pt=28)
        return
    overview = ('%d月%d日，对%s共%d个地市%d个县区%d个镇街%d个村居/监测点开展了蚊媒密度应急监测。'
                '监测结果显示：%d个监测点%s达标准要求，达标率为%.1f%%（%d/%d）；'
                '%d个监测点%s为低中高风险，占%.1f%%（%d/%d）。') % (
        m, d, sec['city_list'], sec['city_num'], sec['district_num'], sec['street_num'],
        sec['point_num'], sec['ok_num'], metric, sec['ok_rate'], sec['ok_num'], sec['total'],
        sec['risk_num'], metric, sec['risk_rate'], sec['risk_num'], sec['total'])
    _add_par(doc, overview, 16, _JUSTIFY, indent_chars=2, line_pt=28)
    for level in ('高风险', '中风险', '低风险'):
        elems = sec['risk_lists'][level]
        if elems:
            _add_par(doc, '%s%d个：%s。' % (level, len(elems), '、'.join(elems)),
                     16, _JUSTIFY, indent_chars=2, line_pt=28)


def write_daily_report(path, target_date, bi_section, adi_section):
    doc = Document()
    _setup_page(doc, 3.70, 3.50, 2.80, 2.60)
    y, m, d = target_date.year, target_date.month, target_date.day

    _add_par(doc, '全省媒介伊蚊传染病相关蚊媒监测情况', 22, _CENTER, line_pt=36)
    _add_par(doc, '', 16, line_pt=15)
    _add_par(doc, '（省疾控中心  %d年%d月%d日20:00）' % (y, m, d), 16, _CENTER, line_pt=28)
    _add_par(doc, '', 16, indent_chars=2, line_pt=28)

    _write_section(doc, '一、布雷图指数（BI）调查评估。', bi_section, m, d)
    _write_section(doc, '二、成蚊密度（ADI）快速评估。', adi_section, m, d)

    doc.save(path)
